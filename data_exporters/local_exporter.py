"""
Data exporter block: write the consolidated dataset.

Decoupled sink: writes the monthly dataset to CSV and to a formatted,
PESTEL-colour-coded multi-sheet Excel workbook (Consolidated Dataset,
Imputation Log, Variable Dictionary, Data Sources). Swapping this block for a
database or cloud writer would not touch any loader or transformer.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from utils.config import COLS, PESTEL_MAP, CSV_OUT, XLSX_OUT, OUTPUTS
from utils.logging_config import get_logger
from data_exporters.workbook_content import IMPUTATION_LOG, VARIABLE_DICTIONARY, DATA_SOURCES

logger = get_logger(__name__)

LIGHT = {"Key": "F2F2F2", "Economic": "EAF6EC", "Social": "DDEEF9",
         "Technological": "FDF1EB", "Note": "FFF9C4"}
HDR = {"Key": "D9D9D9", "Economic": "C6EFCE", "Social": "BDD7EE",
       "Technological": "FCE4D6", "Note": "FFF176"}


def export_csv(df: pd.DataFrame, path: Path = CSV_OUT) -> Path:
    """Write the consolidated monthly dataset to CSV."""
    OUTPUTS.mkdir(parents=True, exist_ok=True)
    df.to_csv(path, index=False)
    logger.info("Wrote CSV: %s (%d rows)", path, len(df))
    return path


def _hdr(ws, row, col, text, bg="1F3864", fg="FFFFFF", wrap=False, sz=9):
    c = ws.cell(row=row, column=col, value=text)
    c.fill = PatternFill("solid", fgColor=bg)
    c.font = Font(bold=True, color=fg, size=sz)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=wrap)


def _cell(ws, row, col, text, bg="FFFFFF", wrap=True, sz=9):
    c = ws.cell(row=row, column=col, value=text)
    c.fill = PatternFill("solid", fgColor=bg)
    c.font = Font(size=sz)
    c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=wrap)


def _write_table_sheet(ws, title, headers, rows, widths, comp_index=None):
    _hdr(ws, 1, 1, title, sz=11)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    for j, h in enumerate(headers, 1):
        _hdr(ws, 2, j, h, wrap=True)
    ws.row_dimensions[2].height = 30
    for ri, row in enumerate(rows, 3):
        comp = row[comp_index] if comp_index is not None else ""
        bg = LIGHT.get(comp, "FFFFFF")
        for ci, val in enumerate(row, 1):
            _cell(ws, ri, ci, val, bg=bg)
        ws.row_dimensions[ri].height = 55
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A3"


def export_workbook(df: pd.DataFrame, path: Path = XLSX_OUT) -> Path:
    """Write the formatted multi-sheet workbook."""
    OUTPUTS.mkdir(parents=True, exist_ok=True)
    all_cols = COLS + ["Imputation_Flag"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Consolidated Dataset"

    ws.append(["PESTEL ->"] + [PESTEL_MAP.get(c, "") for c in all_cols])
    for j, c in enumerate(all_cols, start=2):
        cell = ws.cell(row=1, column=j)
        cell.fill = PatternFill("solid", fgColor=HDR.get(PESTEL_MAP.get(c, ""), "FFFFFF"))
        cell.font = Font(bold=True, size=9)
        cell.alignment = Alignment(horizontal="center")
    ws.append(["#"] + all_cols)
    for j in range(1, len(all_cols) + 2):
        cell = ws.cell(row=2, column=j)
        cell.fill = PatternFill("solid", fgColor="1F3864")
        cell.font = Font(bold=True, color="FFFFFF", size=9)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.row_dimensions[2].height = 35

    imputed = set(df.index[df["Imputation_Flag"] != ""])
    for idx, row in df.iterrows():
        values = [idx + 1]
        for c in all_cols:
            v = row[c]
            if pd.isna(v):
                values.append(None)
            elif c == "Date":
                values.append(pd.to_datetime(v).strftime("%Y-%m-%d"))
            elif isinstance(v, float):
                values.append(round(v, 4))
            else:
                values.append(v)
        ws.append(values)
    for r in range(3, ws.max_row + 1):
        is_imp = (r - 3) in imputed
        for j, c in enumerate(all_cols, start=2):
            comp = PESTEL_MAP.get(c, "")
            cell = ws.cell(row=r, column=j)
            if is_imp and comp in ("Economic", "Technological"):
                cell.fill = PatternFill("solid", fgColor={"Economic": "D5F0D8",
                                        "Technological": "F5D9C8"}.get(comp, "FFFFFF"))
                cell.font = Font(size=9, italic=True)
            else:
                cell.fill = PatternFill("solid", fgColor=LIGHT.get(comp, "FFFFFF"))
                cell.font = Font(size=9)
            cell.alignment = Alignment(horizontal="center")
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 12
    for j, c in enumerate(all_cols[1:], start=3):
        ws.column_dimensions[get_column_letter(j)].width = max(len(c) * 0.85, 11)
    ws.freeze_panes = "C3"

    _write_table_sheet(wb.create_sheet("Imputation Log"),
                       "Imputation Log - PESTEL Monthly Dataset",
                       ["Variable", "PESTEL Component", "Affected Period", "Rows",
                        "Method", "Justification", "Model / Parameters", "Imputed Values"],
                       IMPUTATION_LOG, [26, 16, 18, 8, 14, 38, 38, 42], comp_index=1)
    _write_table_sheet(wb.create_sheet("Variable Dictionary"),
                       "Variable Dictionary - PESTEL Monthly Dataset",
                       ["Column", "PESTEL", "Description", "Unit", "Frequency",
                        "Source", "Coverage", "Imputation"],
                       VARIABLE_DICTIONARY, [26, 14, 42, 20, 14, 30, 18, 26], comp_index=1)
    _write_table_sheet(wb.create_sheet("Data Sources"),
                       "Data Sources - PESTEL Monthly Dataset",
                       ["PESTEL", "Variable(s)", "Source", "Provider",
                        "Series ID", "URL", "Frequency", "Coverage"],
                       DATA_SOURCES, [13, 32, 30, 22, 26, 40, 12, 26], comp_index=0)

    wb.save(path)
    logger.info("Wrote workbook: %s (sheets: %s)", path, [s.title for s in wb.worksheets])
    return path
